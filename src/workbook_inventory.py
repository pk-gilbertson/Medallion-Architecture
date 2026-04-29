"""Workbook inspection utility for baseline candidate workbook."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = REPO_ROOT / "workbook" / "source"
CANDIDATE_FILENAME = "Enterprise_Medallion_Layer_Starter_Kit_Candidate.xlsx"


def inspect_workbook(workbook_path: Path) -> dict[str, Any]:
    """Return workbook inventory metadata for review and change planning."""
    wb = load_workbook(workbook_path, data_only=False)
    inventory: dict[str, Any] = {
        "workbook": str(workbook_path),
        "sheet_names": wb.sheetnames,
        "sheets": [],
        "hidden_sheets": [],
    }

    for ws in wb.worksheets:
        hidden_rows = [idx for idx, dim in ws.row_dimensions.items() if dim.hidden]
        hidden_columns = [str(idx) for idx, dim in ws.column_dimensions.items() if dim.hidden]
        formulas = []
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append(cell.coordinate)

        data_validations = []
        if ws.data_validations is not None:
            for dv in ws.data_validations.dataValidation:
                data_validations.append(
                    {
                        "type": dv.type,
                        "operator": dv.operator,
                        "formula1": dv.formula1,
                        "formula2": dv.formula2,
                        "sqref": str(dv.sqref),
                    }
                )

        tables = []
        for table in ws.tables.values():
            tables.append({"name": table.name, "ref": table.ref})

        merged_ranges = [str(rng) for rng in ws.merged_cells.ranges]
        tab_color = ws.sheet_properties.tabColor.rgb if ws.sheet_properties.tabColor else None

        sheet_summary = {
            "name": ws.title,
            "rows": ws.max_row,
            "columns": ws.max_column,
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "merged_cells": merged_ranges,
            "tables": tables,
            "data_validations": data_validations,
            "formula_cells": formulas,
            "hidden_rows": hidden_rows,
            "hidden_columns": hidden_columns,
            "tab_color": tab_color,
            "sheet_state": ws.sheet_state,
        }

        if ws.sheet_state != "visible":
            inventory["hidden_sheets"].append(ws.title)

        inventory["sheets"].append(sheet_summary)

    wb.close()
    return inventory


def default_candidate_path() -> Path:
    return SOURCE_DIR / CANDIDATE_FILENAME


if __name__ == "__main__":
    candidate_path = default_candidate_path()
    if not candidate_path.exists():
        raise SystemExit(f"Candidate workbook not found at {candidate_path}")

    details = inspect_workbook(candidate_path)
    print(json.dumps(details, indent=2))
